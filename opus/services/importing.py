"""In-memory document extraction and persistence as readable text units."""

from dataclasses import dataclass
from hashlib import sha256
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from posixpath import dirname, normpath
import re
from uuid import uuid4
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from django.db import transaction
from docx import Document
from pypdf import PdfReader

from opus.models import SourceFile, TextUnit


MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 5_000_000
MAX_PARAGRAPHS = 50_000
MAX_EPUB_UNCOMPRESSED_SIZE = 20 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".txt", ".md", ".html", ".htm", ".docx", ".epub", ".pdf"}


class DocumentImportError(ValueError):
    """A safe validation error for a submitted document."""


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    file_type: str
    content_hash: str
    paragraphs: list[str]


class _HTMLTextExtractor(HTMLParser):
    BLOCK_TAGS = {"article", "br", "div", "h1", "h2", "h3", "h4", "h5", "h6", "li", "p", "section"}

    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag.lower() in self.BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag.lower() in self.BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data):
        self.parts.append(data)

    def text(self):
        return "".join(self.parts)


def import_document(edition, uploaded_file):
    """Extract an uploaded document in memory and persist its text units."""

    extracted = extract_document(uploaded_file)
    with transaction.atomic():
        if edition.text_units.exists():
            raise DocumentImportError(
                "This edition already has imported text. Create a new edition to import another document."
            )
        source_file = SourceFile.objects.create(
            version=edition,
            storage_key=f"memory:{edition.pk}:{extracted.content_hash[:16]}:{uuid4().hex[:12]}",
            original_filename=extracted.filename,
            content_hash=extracted.content_hash,
            file_type=extracted.file_type,
            import_status=SourceFile.ImportStatus.COMPLETED,
        )
        TextUnit.objects.bulk_create(
            [
                TextUnit(
                    version=edition,
                    kind=TextUnit.Kind.PARAGRAPH,
                    position=index,
                    content=paragraph,
                )
                for index, paragraph in enumerate(extracted.paragraphs, start=1)
            ]
        )
    return source_file, len(extracted.paragraphs)


def extract_document(uploaded_file):
    """Read a supported upload without ever writing its bytes to disk."""

    filename = Path(uploaded_file.name or "document").name
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentImportError(f"Unsupported file type. Use one of: {allowed}.")
    if uploaded_file.size > MAX_FILE_SIZE:
        raise DocumentImportError("The document must be at most 10 MB.")

    payload = uploaded_file.read()
    if not payload:
        raise DocumentImportError("The uploaded document is empty.")
    text = _extract_text(payload, extension)
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise DocumentImportError("The extracted text is too large to import.")
    paragraphs = _split_paragraphs(text)
    if len(paragraphs) > MAX_PARAGRAPHS:
        raise DocumentImportError("The document contains too many paragraphs to import.")
    if not paragraphs:
        raise DocumentImportError("No readable text was found in the document.")
    return ExtractedDocument(
        filename=filename,
        file_type=extension.removeprefix("."),
        content_hash=sha256(payload).hexdigest(),
        paragraphs=paragraphs,
    )


def _extract_text(payload, extension):
    if extension in {".txt", ".md"}:
        try:
            return payload.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentImportError("Text documents must use UTF-8 encoding.") from exc
    if extension in {".html", ".htm"}:
        try:
            html = payload.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentImportError("HTML documents must use UTF-8 encoding.") from exc
        parser = _HTMLTextExtractor()
        parser.feed(html)
        return parser.text()
    if extension == ".docx":
        try:
            document = Document(BytesIO(payload))
        except Exception as exc:
            raise DocumentImportError("The DOCX document could not be read.") from exc
        return "\n\n".join(paragraph.text for paragraph in document.paragraphs)
    if extension == ".epub":
        return _extract_epub_text(payload)
    if extension == ".pdf":
        try:
            reader = PdfReader(BytesIO(payload))
            if reader.is_encrypted:
                raise DocumentImportError("Password-protected PDFs are not supported.")
            return "\n\n".join(page.extract_text() or "" for page in reader.pages)
        except DocumentImportError:
            raise
        except Exception as exc:
            raise DocumentImportError("The PDF document could not be read.") from exc
    raise DocumentImportError("Unsupported file type.")


def _split_paragraphs(text):
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    blocks = re.split(r"\n\s*\n+", normalized)
    return [" ".join(block.split()) for block in blocks if block.strip()]


def _extract_epub_text(payload):
    try:
        with ZipFile(BytesIO(payload)) as archive:
            entries = archive.infolist()
            if sum(entry.file_size for entry in entries) > MAX_EPUB_UNCOMPRESSED_SIZE:
                raise DocumentImportError("The EPUB expands to too much text to import.")
            container = ElementTree.fromstring(archive.read("META-INF/container.xml"))
            rootfile = next(
                (node.attrib.get("full-path") for node in container.iter() if _local_name(node.tag) == "rootfile"),
                None,
            )
            if not rootfile:
                raise DocumentImportError("The EPUB has no package document.")
            package = ElementTree.fromstring(archive.read(rootfile))
            manifest = {
                node.attrib["id"]: node.attrib["href"]
                for node in package.iter()
                if _local_name(node.tag) == "item" and "id" in node.attrib and "href" in node.attrib
            }
            spine = [
                node.attrib.get("idref")
                for node in package.iter()
                if _local_name(node.tag) == "itemref"
            ]
            chapters = []
            package_directory = dirname(rootfile)
            for item_id in spine:
                href = manifest.get(item_id)
                if not href:
                    continue
                chapter_path = normpath(f"{package_directory}/{href}" if package_directory else href)
                try:
                    chapter = archive.read(chapter_path).decode("utf-8")
                except KeyError as exc:
                    raise DocumentImportError("The EPUB references a missing chapter.") from exc
                except UnicodeDecodeError as exc:
                    raise DocumentImportError("The EPUB contains a chapter that is not UTF-8 encoded.") from exc
                parser = _HTMLTextExtractor()
                parser.feed(chapter)
                chapters.append(parser.text())
            return "\n\n".join(chapters)
    except DocumentImportError:
        raise
    except (BadZipFile, ElementTree.ParseError, KeyError) as exc:
        raise DocumentImportError("The EPUB document could not be read.") from exc


def _local_name(tag):
    return tag.rsplit("}", 1)[-1]
